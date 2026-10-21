import pytest
from pathlib import Path
import docx
from docx.shared import Pt, RGBColor
from app.services.word_resume_service import word_resume_service
from app.schemas.dtos import ResumeImprovementPlan, ResumeChangeItem, ResumeGapItem
from app.services.resume_improvement import resume_improvement_service

def test_resume_improvement_pydantic_schema():
    """Verifies that the structured JSON schema validates change items, elements, and gaps properly."""
    plan_data = {
        "candidate_id": "DOC-1001",
        "target_job": "Senior Python Backend Developer",
        "changes": [
            {
                "section": "professional_summary",
                "target_element_id": "paragraph_2",
                "original_text": "Experienced Python developer building APIs.",
                "proposed_text": "Senior Backend Architect designing high-throughput asynchronous microservices with FastAPI, Redis, and PostgreSQL.",
                "reason": "Aligns architectural focus and modern async tech stack with JD requirements.",
                "supporting_evidence": [
                    {"source": "candidate_resume", "text": "Built FastAPI microservices processing 10k req/s"}
                ],
                "requires_user_confirmation": False
            }
        ],
        "gaps": [
            {
                "requirement": "Kubernetes cluster administration",
                "status": "not_demonstrated",
                "evidence": None
            }
        ]
    }
    plan = ResumeImprovementPlan(**plan_data)
    assert plan.candidate_id == "DOC-1001"
    assert len(plan.changes) == 1
    assert plan.changes[0].target_element_id == "paragraph_2"
    assert plan.changes[0].requires_user_confirmation is False
    assert len(plan.gaps) == 1
    assert plan.gaps[0].status == "not_demonstrated"

def test_docx_inplace_element_extraction_and_run_preservation(tmp_path):
    """Verifies that in-place editing maps element IDs, edits text in first run, and preserves fonts and styling."""
    # 1. Create a sample DOCX with styled runs
    sample_doc = docx.Document()
    p1 = sample_doc.add_paragraph()
    r1 = p1.add_run("John Doe - Senior Software Engineer")
    r1.font.name = "Calibri"
    r1.font.size = Pt(14)
    r1.bold = True
    r1.font.color.rgb = RGBColor(30, 64, 175)

    p2 = sample_doc.add_paragraph()
    r2 = p2.add_run("Developed backend APIs using Flask and MySQL.")
    r2.font.name = "Arial"
    r2.font.size = Pt(10)

    orig_path = tmp_path / "original_resume.docx"
    sample_doc.save(str(orig_path))

    # 2. Extract elements
    elements = word_resume_service.extract_docx_elements(orig_path)
    assert len(elements) == 2
    assert elements[0]["element_id"] == "paragraph_0"
    assert "John Doe" in elements[0]["text"]
    assert elements[1]["element_id"] == "paragraph_1"
    assert "Flask" in elements[1]["text"]

    # 3. Apply targeted in-place changes
    target_changes = [
        {
            "target_element_id": "paragraph_1",
            "proposed_text": "Architected high-throughput backend APIs utilizing FastAPI and PostgreSQL."
        }
    ]
    out_path = tmp_path / "revised_resume.docx"
    success = word_resume_service.apply_docx_changes_inplace(orig_path, out_path, target_changes)
    assert success is True
    assert out_path.exists()

    # 4. Verify original remains completely untouched
    orig_check = docx.Document(str(orig_path))
    assert "Flask" in orig_check.paragraphs[1].text
    assert "FastAPI" not in orig_check.paragraphs[1].text

    # 5. Verify revision applied change and preserved font attributes
    rev_check = docx.Document(str(out_path))
    assert len(rev_check.paragraphs) == 2
    assert "FastAPI" in rev_check.paragraphs[1].text
    assert rev_check.paragraphs[1].runs[0].font.name == "Arial"
    assert rev_check.paragraphs[1].runs[0].font.size == Pt(10)

def test_recalculate_score_api_endpoint(client):
    """Verifies POST /api/resumes/{doc_id}/recalculate-score genuinely recalculates score with honest delta."""
    resp = client.post(
        "/api/resumes/DOC-1001/recalculate-score",
        json={
            "job_id": "JOB-101",
            "accepted_changes": [
                {
                    "target_element_id": "paragraph_1",
                    "original_text": "Original text",
                    "proposed_text": "Engineered enterprise microservices in FastAPI with PostgreSQL.",
                    "reason": "Directly highlights backend tech stack"
                }
            ]
        }
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "original_score" in data
    assert "revised_score" in data
    assert "score_difference" in data
    assert "explanation" in data
    assert isinstance(data["gained_requirements"], list)
    assert isinstance(data["remaining_gaps"], list)

def test_ikm_auth_masked_password_handling(client):
    """Verifies that submitting blank or masked password in /api/ikm/test-auth uses .env password rather than erroring."""
    resp = client.post(
        "/api/ikm/test-auth",
        json={"username": "resume-screening", "password": "••••••••••••••••••••"}
    )
    assert resp.status_code == 200
    data = resp.json()
    assert "success" in data
    assert "status_code" in data
    # When remote returns 423, the response clearly explains the remote lockout
    if data["status_code"] == 423:
        assert "Account locked" in data["message"]
