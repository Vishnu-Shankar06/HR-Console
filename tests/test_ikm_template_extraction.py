import io
import re
from datetime import datetime
import pytest
from unittest.mock import patch, MagicMock

from app.integrations.ikm_adapter import (
    HttpIKMAdapter,
    LocalIKMAdapter,
    normalize_template_data
)
from app.services.experience_calculator import experience_calculator
from app.integrations.grok_adapter import mask_candidate_pii, restore_candidate_pii
from app.services.word_resume_service import word_resume_service
from app.schemas.dtos import CandidatePIIDTO, CandidateProfileDTO


# =====================================================================
# 1. IKM Template-Extraction Endpoints & Polling Tests
# =====================================================================

def test_post_start_template_extraction():
    """Verify POST /api/v1/documents/{document_id}/template-extraction initiates extraction."""
    adapter = HttpIKMAdapter(
        base_url="http://test-ikm.local",
        username="user",
        password="pass",
        timeout=5.0
    )
    with patch.object(adapter, "authenticate", return_value=True), \
         patch.object(adapter._client, "post") as mock_post:
        mock_post.return_value = MagicMock(
            status_code=200,
            json=lambda: {
                "document_id": "doc-uuid-1234",
                "queued": True,
                "processor_type": "template_extraction",
                "status": "QUEUED",
                "message": "Template extraction queued."
            }
        )
        res = adapter.start_template_extraction("doc-uuid-1234")
        assert res.get("status") == "QUEUED"
        assert res.get("queued") is True
        assert res.get("document_id") == "doc-uuid-1234"
        mock_post.assert_called_once()
        assert "/api/v1/documents/doc-uuid-1234/template-extraction" in mock_post.call_args[0][0]


def test_get_stored_template_and_alias():
    """Verify GET /documents/{id}/template and GET /documents/{id}/template-extraction."""
    adapter = HttpIKMAdapter(
        base_url="http://test-ikm.local",
        username="user",
        password="pass",
        timeout=5.0
    )
    sample_tpl = {
        "document_id": "doc-uuid-1234",
        "contact_block": {"name": "Alex Smith", "email": "alex@example.com"},
        "target_role": "Backend Engineer",
        "skills": ["Python", "FastAPI"]
    }
    with patch.object(adapter, "authenticate", return_value=True), \
         patch.object(adapter._client, "get") as mock_get:
        mock_get.return_value = MagicMock(status_code=200, json=lambda: sample_tpl)
        
        # Test GET /template
        res1 = adapter.get_stored_template("doc-uuid-1234")
        assert res1 is not None
        assert res1["contact_block"]["name"] == "Alex Smith"
        assert "/api/v1/documents/doc-uuid-1234/template" in mock_get.call_args[0][0]

        # Test GET /template-extraction (alias)
        res2 = adapter.get_template_extraction_result("doc-uuid-1234")
        assert res2 is not None
        assert res2["target_role"] == "Backend Engineer"


def test_template_polling_with_bounded_retries():
    """Verify bounded polling retries and succeeds when template becomes ready."""
    adapter = HttpIKMAdapter(
        base_url="http://test-ikm.local",
        username="user",
        password="pass",
        timeout=5.0
    )
    with patch.object(adapter, "authenticate", return_value=True), \
         patch.object(adapter, "get_stored_template", side_effect=[None, None, {"status": "ready", "skills": ["Python"]}]) as mock_get:
        res = adapter.poll_template_extraction("doc-uuid-1234", max_retries=3, retry_delay=0.01)
        assert res is not None
        assert res["skills"] == ["Python"]
        assert mock_get.call_count == 3


def test_template_polling_incomplete_graceful():
    """Verify polling stops after max retries without hanging when processing is incomplete."""
    adapter = HttpIKMAdapter(
        base_url="http://test-ikm.local",
        username="user",
        password="pass",
        timeout=5.0
    )
    with patch.object(adapter, "authenticate", return_value=True), \
         patch.object(adapter, "get_stored_template", return_value=None), \
         patch.object(adapter, "get_template_extraction_result", return_value=None):
        res = adapter.poll_template_extraction("doc-uuid-1234", max_retries=2, retry_delay=0.01)
        assert res is None


def test_fallback_to_local_parsing_on_404():
    """Verify fetch_or_extract_template falls back seamlessly to local parser on 404."""
    adapter = HttpIKMAdapter(
        base_url="http://test-ikm.local",
        username="user",
        password="pass",
        timeout=5.0
    )
    raw_text = (
        "John Doe\n"
        "Email: john.doe@example.com\n"
        "Phone: +1 555-0199\n"
        "Location: Chicago, IL\n"
        "PROFESSIONAL SUMMARY\n"
        "Experienced Software Engineer specializing in Python and cloud services.\n"
        "WORK EXPERIENCE\n"
        "Senior Developer | Tech Corp | June 2021 – Present\n"
        "- Developed high-throughput microservices using FastAPI and PostgreSQL.\n"
    )
    with patch.object(adapter, "authenticate", return_value=True), \
         patch.object(adapter, "get_stored_template", return_value=None), \
         patch.object(adapter, "poll_template_extraction", return_value=None):
        res = adapter.fetch_or_extract_template("doc-uuid-404", raw_text=raw_text)
        assert res is not None
        assert res["name"] == "John Doe"
        assert res["email"] == "john.doe@example.com"
        assert res["experience_years"] > 0
        assert "yrs" in res["experience_display"] or "mos" in res["experience_display"]


# =====================================================================
# 2. Metadata Normalization Schema Tests
# =====================================================================

def test_metadata_normalization_various_shapes():
    """Verify normalize_template_data handles diverse nested and flat shapes."""
    # Shape 1: Nested under 'document_template'
    shape1 = {
        "document_template": {
            "contact_block": {
                "name": "Sarah Connor",
                "email": "sarah@resistance.org",
                "phone": "+1 310-555-0100",
                "location": "Los Angeles, CA",
                "links": "https://linkedin.com/in/sconnor | https://github.com/sconnor"
            },
            "headline": "DevOps Architect",
            "summary": "Specialist in high-availability distributed systems.",
            "skills": ["Kubernetes", "Docker", "Terraform", "AWS"],
            "experience": [
                {
                    "company": "Cyberdyne Systems",
                    "role": "Systems Engineer",
                    "dates": "January 2021 - December 2023",
                    "bullets": ["Managed Kubernetes clusters", "Automated CI/CD pipelines"]
                }
            ],
            "education": [{"degree": "B.S. Computer Engineering", "institution": "MIT", "graduation_year": "2020"}],
            "projects": [{"title": "Defense Matrix", "description": "Cloud security system", "technologies": ["AWS", "Terraform"]}],
            "certifications": [{"name": "AWS Certified DevOps Pro", "issuing_body": "Amazon", "date": "2022"}]
        }
    }
    norm1 = normalize_template_data(shape1)
    assert norm1["name"] == "Sarah Connor"
    assert norm1["title"] == "DevOps Architect"
    assert norm1["email"] == "sarah@resistance.org"
    assert norm1["phone"] == "+1 310-555-0100"
    assert norm1["location"] == "Los Angeles, CA"
    assert norm1["linkedin"] == "https://linkedin.com/in/sconnor"
    assert norm1["github"] == "https://github.com/sconnor"
    assert "Kubernetes" in norm1["skills"]
    assert 2.8 <= norm1["experience_years"] <= 3.0
    assert "yr" in norm1["experience_display"]
    assert len(norm1["education"]) == 1
    assert len(norm1["projects"]) == 1
    assert len(norm1["certifications"]) == 1


# =====================================================================
# 3. Deterministic Experience Calculation Tests
# =====================================================================

def test_experience_calculation_diverse_date_formats():
    """Test experience calculation across standard and edge-case date formats."""
    now = datetime(2026, 10, 1)

    # 1. Month Year - Present
    p1 = experience_calculator.parse_date_range("June 2023 – Present", now=now)
    assert p1["is_valid"] is True
    assert p1["months"] == (2026 - 2023) * 12 + (10 - 6)  # 40 months

    # 2. MM/YYYY - MM/YYYY
    p2 = experience_calculator.parse_date_range("01/2020 – 12/2022", now=now)
    assert p2["is_valid"] is True
    assert p2["months"] == (2022 - 2020) * 12 + (12 - 1)  # 35 months

    # 3. YYYY - YYYY
    p3 = experience_calculator.parse_date_range("2019 – 2021", now=now)
    assert p3["is_valid"] is True
    assert p3["months"] == 24

    # 4. YYYY-MM to YYYY-MM
    p4 = experience_calculator.parse_date_range("2022-03 to 2023-08", now=now)
    assert p4["is_valid"] is True
    assert p4["months"] == (2023 - 2022) * 12 + (8 - 3)  # 17 months


def test_experience_calculation_overlapping_intervals():
    """Verify concurrent/overlapping roles are merged and NOT double-counted."""
    now = datetime(2026, 10, 1)
    roles = [
        {"company": "Company A", "role": "Senior Engineer", "dates": "01/2020 – 12/2021"},
        {"company": "Company B (Concurrent)", "role": "Contractor", "dates": "06/2020 – 06/2021"},
        {"company": "Company C", "role": "Lead Architect", "dates": "01/2022 – 12/2023"}
    ]
    calc = experience_calculator.calculate_total_experience(roles, now=now)
    # Company A: 01/2020 to 12/2021 = 23 months
    # Company B: 06/2020 to 06/2021 is entirely inside Company A interval
    # Company C: 01/2022 to 12/2023 = 23 months
    # Total non-overlapping = 23 + 23 = 46 months (~3.8 yrs), NOT 46 + 12 = 58 months!
    assert calc["total_months"] == 46
    assert calc["total_years"] == 3.8
    assert "yrs" in calc["display"]


def test_experience_calculation_invalid_dates():
    """Verify end date preceding start date is marked invalid and excluded."""
    p_invalid = experience_calculator.parse_date_range("2024 – 2020")
    assert p_invalid["is_valid"] is False
    assert p_invalid["months"] == 0


def test_experience_display_no_object_object():
    """Verify experience display NEVER produces [object Object], NaN, or null yrs."""
    empty_calc = experience_calculator.calculate_total_experience([])
    assert empty_calc["total_years"] == 0.0
    assert empty_calc["display"] == "0 mos"
    assert "[object" not in empty_calc["display"]
    assert "NaN" not in empty_calc["display"]
    assert "null" not in empty_calc["display"]


# =====================================================================
# 4. Privacy Boundary & Groq Anonymization Tests
# =====================================================================

def test_privacy_preserving_groq_masking():
    """Verify candidate personal info is masked before transmission to LLM."""
    real_text = (
        "Jonathan Wick\n"
        "Email: j.wick@continental.com\n"
        "Phone: +1 212-555-0199\n"
        "Location: New York, NY\n"
        "LinkedIn: https://linkedin.com/in/jonathanwick\n"
        "Expert in systems automation and backend architecture.\n"
    )
    pd = {
        "name": "Jonathan Wick",
        "email": "j.wick@continental.com",
        "phone": "+1 212-555-0199",
        "location": "New York, NY",
        "links": "https://linkedin.com/in/jonathanwick"
    }
    masked, token_map = mask_candidate_pii(real_text, pd)

    # Real personal data must NOT appear in the masked text
    assert "Jonathan Wick" not in masked
    assert "j.wick@continental.com" not in masked
    assert "+1 212-555-0199" not in masked
    assert "New York, NY" not in masked
    assert "https://linkedin.com/in/jonathanwick" not in masked

    # Placeholders must be present
    assert "[CANDIDATE_NAME]" in masked
    assert "[CANDIDATE_EMAIL]" in masked or "[EMAIL]" in masked
    assert "[CANDIDATE_PHONE]" in masked or "[PHONE]" in masked

    # Restoration must bring back original details
    restored = restore_candidate_pii(masked, token_map)
    assert "Jonathan Wick" in restored
    assert "j.wick@continental.com" in restored


# =====================================================================
# 5. Word DOCX Restoration & Personal Overrides Tests
# =====================================================================

def test_docx_personal_details_restoration():
    """Verify word_resume_service updates contact details without destroying layout."""
    from docx import Document as DocxDocument
    doc = DocxDocument()
    doc.add_paragraph("OLD CANDIDATE NAME")
    doc.add_paragraph("old.email@test.com | +1 000-000-0000 | Old City, TX")
    doc.add_paragraph("PROFESSIONAL SUMMARY")
    doc.add_paragraph("Experienced engineer.")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    overrides = {
        "name": "New Candidate",
        "email": "new.candidate@example.com",
        "phone": "+1 555-9876",
        "location": "Boston, MA",
        "role": "Lead Architect"
    }

    out_bytes = word_resume_service.apply_docx_changes_inplace(
        original_source=docx_bytes,
        approved_changes=[],
        personal_overrides=overrides,
        output_dest=None
    )
    assert out_bytes is not None
    assert len(out_bytes) > 0

    # Read back modified document
    modified_doc = DocxDocument(io.BytesIO(out_bytes))
    full_text = "\n".join(p.text for p in modified_doc.paragraphs)
    assert "NEW CANDIDATE" in full_text.upper()
    assert "new.candidate@example.com" in full_text
    assert "+1 555-9876" in full_text
    assert "[CANDIDATE_NAME]" not in full_text


# =====================================================================
# 6. DTO Schema Validation Tests
# =====================================================================

def test_candidate_pii_dto_schema():
    """Verify CandidatePIIDTO contains all authorized recruiter view fields."""
    dto = CandidatePIIDTO(
        doc_id="DOC-1001",
        name="Elena Rostova",
        title="Full Stack Engineer",
        role="Full Stack Engineer",
        email="elena@example.com",
        phone="+1 415-555-0144",
        address="San Francisco, CA",
        location="San Francisco, CA",
        links="https://linkedin.com/in/elena",
        linkedin="https://linkedin.com/in/elena",
        filename="Elena_Resume.docx",
        uploaded_at="2026-10-29T10:00:00Z",
        experience_years=4.5,
        experience_display="4 yrs 6 mos",
        skills=["React", "TypeScript", "Python"]
    )
    assert dto.name == "Elena Rostova"
    assert dto.title == "Full Stack Engineer"
    assert dto.experience_years == 4.5
    assert dto.experience_display == "4 yrs 6 mos"
    assert "[object" not in dto.experience_display
