from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document

BASE = Path(__file__).resolve().parent.parent
SAMPLES_DIR = BASE / "sample_data" / "resumes"
JOBS_DIR = BASE / "sample_data" / "jobs"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
JOBS_DIR.mkdir(parents=True, exist_ok=True)

def create_pdf(filename: str, name: str, email: str, role: str, skills: list, summary: str):
    file_path = SAMPLES_DIR / filename
    c = canvas.Canvas(str(file_path), pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, 750, name)
    c.setFont("Helvetica", 10)
    c.drawString(50, 735, f"Email: {email} | Phone: +91-9876543210 | Location: Bengaluru, India")
    c.setFont("Helvetica-Bold", 12)
    c.drawString(50, 705, f"Role: {role}")
    c.setFont("Helvetica", 10)
    c.drawString(50, 685, f"Skills: {', '.join(skills)}")
    c.drawString(50, 660, "Professional Summary & Projects:")
    
    text_obj = c.beginText(50, 640)
    text_obj.setFont("Helvetica", 9)
    for line in summary.split("\n"):
        text_obj.textLine(line.strip())
    c.drawText(text_obj)
    c.save()
    print(f"Created PDF: {filename}")

def create_docx(filename: str, name: str, email: str, role: str, skills: list, summary: str):
    file_path = SAMPLES_DIR / filename
    doc = Document()
    doc.add_heading(name, level=0)
    doc.add_paragraph(f"Email: {email} | Phone: +91-9876501234 | Location: Pune, India")
    doc.add_heading(f"Role: {role}", level=1)
    doc.add_paragraph(f"Skills: {', '.join(skills)}")
    doc.add_heading("Projects & Experience", level=2)
    doc.add_paragraph(summary)
    doc.save(str(file_path))
    print(f"Created DOCX: {filename}")

# Generate samples
create_pdf(
    "resume_arun_backend.pdf",
    "Arun Kumar",
    "arun.kumar.dev@example.com",
    "Backend Developer",
    ["Python", "FastAPI", "PostgreSQL", "REST APIs", "SQLAlchemy", "Docker", "Git"],
    """Project: High-Throughput E-Commerce Payment Gateway
- Developed REST API endpoints using FastAPI and PostgreSQL handling 5M requests/day.
- Integrated SQLAlchemy ORM with connection pooling and PostgreSQL indexing.
- Containerized microservices using Docker and orchestrated multi-service testing with PyTest."""
)

create_docx(
    "resume_priya_fullstack.docx",
    "Priya Sharma",
    "priya.sharma.tech@example.com",
    "Full Stack Engineer",
    ["Python", "FastAPI", "PostgreSQL", "React", "TypeScript", "REST APIs", "TailwindCSS"],
    """Project: Teleconsultation Healthcare Portal
- Designed interactive patient scheduling portal using React, TypeScript, and modern CSS.
- Built backend REST APIs with Python FastAPI and PostgreSQL database tables.
- Implemented real-time booking reminders and automated appointment workflows."""
)

create_pdf(
    "resume_vikram_devops.pdf",
    "Vikram Patel",
    "vikram.patel.cloud@example.com",
    "DevOps Engineer",
    ["Python", "Docker", "Kubernetes", "AWS", "CI/CD", "Linux", "Git"],
    """Project: Cloud Automation & Deployment Pipeline
- Architected automated GitHub Actions CI/CD pipelines deploying to AWS ECS.
- Configured Kubernetes clusters and container management policies.
- Scripted operational automation tools using Python and Linux shell scripts."""
)

create_pdf(
    "resume_ananya_data.pdf",
    "Ananya Sen",
    "ananya.sen.ai@example.com",
    "Data Scientist",
    ["Python", "Pandas", "NumPy", "SQL", "PostgreSQL", "PyTorch", "FastAPI"],
    """Project: Churn Prediction & Machine Learning Services
- Analyzed enterprise customer datasets using Pandas, NumPy, and SQL queries.
- Built predictive models and deployed inference microservices using FastAPI.
- Validated performance metrics and stored analytics predictions in PostgreSQL."""
)

create_docx(
    "resume_rohit_frontend.docx",
    "Rohit Gupta",
    "rohit.gupta.fe@example.com",
    "Frontend Developer",
    ["React", "TypeScript", "JavaScript", "Next.js", "HTML5", "CSS3", "Git"],
    """Project: Financial Insights Dashboard
- Engineered performant interactive dashboards with React, Next.js, and TypeScript.
- Integrated charting libraries with live REST endpoints and responsive layouts."""
)

# Sample JD files
with open(JOBS_DIR / "senior_python_backend.txt", "w", encoding="utf-8") as f:
    f.write("""Job Title: Senior Python Backend Developer
Location: Remote
Requirements:
- Python proficiency with FastAPI or Django
- PostgreSQL database experience and query optimization
- REST API design and microservices
- Docker containerization and PyTest test suites
- 2+ years backend engineering experience
""")

print("All sample files generated successfully!")
